import openpyxl

wb = openpyxl.load_workbook(r'c:\Users\Pradeep S\Desktop\Semester VI\Project Phase - 1\Code\benchmark_winrates.xlsx')
ws = wb.active

headers = [cell.value for cell in ws[1]]

# Get all algorithm data
algos = {}
for row in ws.iter_rows(min_row=2, values_only=False):
    name = row[0].value
    if not name:
        continue
    vals = [row[i].value for i in range(1, len(headers))]
    algos[name] = vals

# Runs 4-7 (0-indexed pairs: index 6,7 = Run4,Time4 etc.)
# Run4 starts at index 6, Run5 at 8, Run6 at 10, Run7 at 12
run_indices = [6, 8, 10, 12]  # indices for Run4, Run5, Run6, Run7
run_labels = ["Run 4", "Run 5", "Run 6", "Run 7"]

print("=" * 70)
print("ABLATION ANALYSIS for Runs 4-7 (missing from ablation_results.txt)")
print("=" * 70)

ablation_text = []

for idx, ri in enumerate(run_indices):
    cce = algos["CCE-TDCCEA"][ri]
    std = algos["Standard CCEA"][ri]
    td = algos["TD-CCEA"][ri]
    gesy = algos["GESY"][ri]
    pdcoea = algos["SA-PDCoEA"][ri]
    
    print("\n--- {} (CCE-TDCCEA = {:.1f}%) ---".format(run_labels[idx], cce))
    
    block = []
    block.append("\nRun Date: {} (from xlsx)".format(run_labels[idx]))
    
    # Trajectory
    d = cce - std
    if d >= 0:
        line1 = "\n  Trajectory memory contribution:"
        line2 = "    CCE-TDCCEA ({:.1f}%) vs Std-CCEA ({:.1f}%)".format(cce, std)
        line3 = "    Result: Trajectory HELPS (+{:.1f}%)".format(abs(d))
    else:
        line1 = "\n  Trajectory memory contribution:"
        line2 = "    CCE-TDCCEA ({:.1f}%) vs Std-CCEA ({:.1f}%)".format(cce, std)
        line3 = "    Result: Trajectory HURTS ({:.1f}%)".format(abs(d))
    block.extend([line1, line2, line3])
    print(line1); print(line2); print(line3)
    
    # Decomposition
    d = cce - td
    if d >= 0:
        line1 = "\n  Cooperative decomposition contribution:"
        line2 = "    CCE-TDCCEA ({:.1f}%) vs TD-CCEA ({:.1f}%)".format(cce, td)
        line3 = "    Result: Decomposition HELPS (+{:.1f}%)".format(abs(d))
    else:
        line1 = "\n  Cooperative decomposition contribution:"
        line2 = "    CCE-TDCCEA ({:.1f}%) vs TD-CCEA ({:.1f}%)".format(cce, td)
        line3 = "    Result: Decomposition HURTS ({:.1f}%)".format(abs(d))
    block.extend([line1, line2, line3])
    print(line1); print(line2); print(line3)
    
    # GESY
    d = gesy - cce
    if d >= 0:
        line1 = "\n  GESY seeded GP contribution:"
        line2 = "    GESY ({:.1f}%) vs CCE-TDCCEA ({:.1f}%)".format(gesy, cce)
        line3 = "    Result: GESY BETTER (+{:.1f}%)".format(abs(d))
    else:
        line1 = "\n  GESY seeded GP contribution:"
        line2 = "    GESY ({:.1f}%) vs CCE-TDCCEA ({:.1f}%)".format(gesy, cce)
        line3 = "    Result: GESY WORSE ({:.1f}%)".format(abs(d))
    block.extend([line1, line2, line3])
    print(line1); print(line2); print(line3)
    
    # SA-PDCoEA
    d = pdcoea - cce
    if d >= 0:
        line1 = "\n  SA-PDCoEA pairwise dominance contribution:"
        line2 = "    SA-PDCoEA ({:.1f}%) vs CCE-TDCCEA ({:.1f}%)".format(pdcoea, cce)
        line3 = "    Pairwise dominance selection: SA-PDCoEA BETTER (+{:.1f}%)".format(abs(d))
    else:
        line1 = "\n  SA-PDCoEA pairwise dominance contribution:"
        line2 = "    SA-PDCoEA ({:.1f}%) vs CCE-TDCCEA ({:.1f}%)".format(pdcoea, cce)
        line3 = "    Pairwise dominance selection: SA-PDCoEA WORSE ({:.1f}%)".format(abs(d))
    block.extend([line1, line2, line3])
    print(line1); print(line2); print(line3)
    
    ablation_text.append("\n".join(block))

print("\n" + "=" * 70)
print("NOTE: Wilcoxon test CANNOT be computed for these 4 runs because")
print("per-bot win rates are not available (only overall % from xlsx).")
print("=" * 70)

# Output text ready to append
print("\n\n--- TEXT TO APPEND TO ablation_results.txt ---\n")
for t in ablation_text:
    print(t)
    print()
