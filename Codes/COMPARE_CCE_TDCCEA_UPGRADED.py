#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
COMPARE_CCE_TDCCEA_UPGRADED.py
==============================
Strict Head-to-Head Comparison:
  1. CCE-TDCCEA (Baseline)
     — Standard Cooperative Co-Evolutionary TD-CCEA (matching COMPARE_ALL.py):
       • Macro (8) and Micro (12) sub-populations
       • Top-3 partner cooperative evaluation
       • Coordinated environment centroid (ca + cb) with sub-agent trajectory memory
       • Redundancy pruning (delta_s=0.85) and fitness filtering (delta_f=0.20)
       • Sub-population elite archive (A_max=15) and tournament selection (k=3)
  2. CCE-TDCCEA (Upgraded Trajectory)
     — The exact same Cooperative Co-Evolutionary framework, but upgraded with the
       trajectory mechanism from "TDDCEA_Modification.pdf":
       • Clustered strategic modes fingerprint F_t = {(theta_m, w_m, g_m)} per sub-population
       • Optimal-Transport (Earth-Mover) environmental distance D_env
       • Mode-weighted similarity trajectory scoring: alpha_k = max(0, 1 - D_k / D_max)

EVERY OTHER COMPONENT IS IDENTICAL:
  - Same problem: 20-D strategy weights for microRTS.
  - Same sub-populations: Macro (8 weights) and Micro (12 weights).
  - Same partner collaboration: top-3 cooperative evaluation.
  - Same population sizes: N=20 for Macro, N=20 for Micro.
  - Same selection: Redundancy filter (delta_s = 0.85), weak filter (delta_f = 0.20).
  - Same archive: Elite Hall of Fame (top 3, cap A_max = 15).
  - Same reproduction: Tournament selection (k=3) from survivors U archive with clamped mutations.
  - Same opponent rotation: 7 built-in bots across 30 generations.
  - Multi-core parallelism: ServerPool with parallel Java game runners.

OUTPUTS:
  - Per-opponent breakdown table.
  - Wilcoxon signed-rank test (statistical significance).
  - Ablation analysis (impact of the upgraded trajectory mechanism).
  - Convergence table (generation-by-generation best combined fitness).
  - JSON & Excel exports.
"""

import math
import os
import sys
import json
import time
import queue
import random
import subprocess
from concurrent.futures import ThreadPoolExecutor

import numpy as np

# Force UTF-8 encoding on Windows console
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

try:
    from scipy.optimize import linear_sum_assignment as _lsa
except Exception:
    _lsa = None

try:
    import openpyxl
    _OPENPYXL_AVAILABLE = True
except ImportError:
    _OPENPYXL_AVAILABLE = False

os.environ["PYTHONIOENCODING"] = "utf-8"

# ===============================================================
#  CONFIGURATION & PATH DISCOVERY
# ===============================================================
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

cand_dirs = [
    os.path.join(os.path.dirname(os.path.dirname(SCRIPT_DIR)), "Simulation", "microrts"),
    os.path.join(os.path.dirname(SCRIPT_DIR), "Simulation", "microrts"),
    os.path.join(SCRIPT_DIR, "Simulation", "microrts"),
]
MICRORTS_DIR = next(
    (d for d in cand_dirs if os.path.isfile(os.path.join(d, "bin", "tests", "EvolvedBatchGameRunner.class"))),
    cand_dirs[0]
)

# Standard arena chambers 32x32 map used in COMPARE_ALL.py
MAP_PATH = "maps/arena_chambers_32x32.xml"
if not os.path.isfile(os.path.join(MICRORTS_DIR, MAP_PATH)):
    for _fb in ["maps/chambers32x32.xml", "maps/basesWorkers32x32A.xml"]:
        if os.path.isfile(os.path.join(MICRORTS_DIR, _fb)):
            MAP_PATH = _fb
            break

MAX_CYCLES = 10000
JAVA_CMD   = "java"

WEIGHT_NAMES = [
    "worker_ratio", "max_workers", "expand_threshold", "light_preference",
    "heavy_preference", "ranged_preference", "barracks_timing", "aggression",
    "attack_threshold", "defend_radius", "target_priority", "defense_trigger",
    "counter_rush", "tech_switch", "economy_pressure", "scout_priority",
    "group_attack", "second_barracks", "worker_fight_ratio", "focus_fire",
]
DIMENSIONS = len(WEIGHT_NAMES)  # 20
D = DIMENSIONS

# Sub-population indices (Cooperative Co-evolutionary decomposition)
MACRO_INDICES = [0, 1, 2, 6, 14, 15, 17, 18]                  # 8 weights (Economy & Macro)
MICRO_INDICES = [3, 4, 5, 7, 8, 9, 10, 11, 12, 13, 16, 19]   # 12 weights (Combat & Micro)

OPPONENT_ROTATION = [
    "EconomyMilitaryRush", "LightRush", "WorkerRush",
    "HeavyDefense", "RangedRush", "EconomyRush", "WorkerRushPlusPlus",
]
N_OPP = len(OPPONENT_ROTATION)

# Shared benchmark hyperparameters (from COMPARE_ALL.py)
N_POP             = 20      # Population size per sub-population
K_GAMES           = 3       # Games per evaluation
TOTAL_GENERATIONS = 30      # Generations
TRAJ_WINDOW       = 7       # History window W
MU_BASE           = 0.15    # Base mutation scale
EPS               = 0.03    # Mutation offset
DELTA_F           = 0.20    # Weak-agent fitness threshold
DELTA_S           = 0.85    # Redundancy similarity threshold
A_MAX             = 15      # Elite archive capacity per sub-population
TOP_PARTNERS      = 3       # Collaborating partners in cooperative evaluation
VALIDATION_GAMES  = 5       # Validation games per opponent
DRAW_SCORE        = 0.3     # Draw partial reward
MIN_KEEP          = 8       # Minimum survivors floor

# Upgraded trajectory hyperparameters (TDDCEA_Modification.pdf)
TAU0  = 0.15   # Base split threshold
LAM   = 0.50   # Threshold growth: tau_M = TAU0 * exp(LAM * (M - 1))
D_MAX = 0.40   # Distance cutoff for similarity

_fp_counter = 0


def next_fp_uid():
    global _fp_counter
    _fp_counter += 1
    return _fp_counter


def clamp01(x):
    return max(0.0, min(1.0, float(x)))


def combine_subagents(macro, micro):
    """Combine macro and micro sub-agent parameters into a full 20-D weight vector."""
    full = [0.5] * DIMENSIONS
    macro_vals = macro.params if hasattr(macro, "params") else macro
    micro_vals = micro.params if hasattr(micro, "params") else micro
    for i, idx in enumerate(MACRO_INDICES):
        full[idx] = clamp01(macro_vals[i])
    for i, idx in enumerate(MICRO_INDICES):
        full[idx] = clamp01(micro_vals[i])
    return full


# ===============================================================
#  MULTI-CORE GAME POOL & SERVER
# ===============================================================
def default_workers():
    try:
        cores = len(os.sched_getaffinity(0))
    except AttributeError:
        cores = os.cpu_count() or 4
    env = os.environ.get("CCE_WORKERS")
    if env:
        try:
            return max(1, int(env))
        except ValueError:
            pass
    return max(1, min(cores - 1, N_POP))


class GameServer:
    """Manages an individual microRTS Java process via standard I/O pipes."""

    def __init__(self):
        self.process = None

    def start(self):
        classpath = os.pathsep.join(["lib/*", "bin"])
        self.process = subprocess.Popen(
            [JAVA_CMD, "-cp", classpath, "tests.EvolvedBatchGameRunner"],
            cwd=MICRORTS_DIR,
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            encoding="utf-8",
            bufsize=1,
        )
        line = self.process.stdout.readline().strip()
        if line != "READY":
            err = ""
            try:
                err = self.process.stderr.read(1000)
            except Exception:
                pass
            raise RuntimeError(f"Server failed to start: stdout='{line}', stderr='{err}'")

    def run_game(self, weights, opponent):
        ws = ";".join(f"{w:.4f}" for w in weights)
        req = f"EVOLVED:{ws},{opponent},{MAP_PATH},{MAX_CYCLES}\n"
        try:
            self.process.stdin.write(req)
            self.process.stdin.flush()
            line = self.process.stdout.readline().strip()
            if line.startswith("RESULT:"):
                p = line.replace("RESULT:", "").split(",")
                return int(p[0]), int(p[1])
        except (BrokenPipeError, OSError):
            pass
        return -1, MAX_CYCLES

    def stop(self):
        if self.process:
            try:
                self.process.stdin.write("EXIT\n")
                self.process.stdin.flush()
                self.process.wait(timeout=5)
            except Exception:
                self.process.kill()
            finally:
                self.process = None


class ServerPool:
    """Pool of GameServer processes; game evaluations are dispatched to idle workers."""

    def __init__(self, base_server=None, n_workers=None):
        self.n = n_workers or default_workers()
        self.base = base_server
        self.owned_servers = []
        self.q = queue.Queue()

        if base_server is not None:
            self.q.put(base_server)
            extra_count = self.n - 1
        else:
            extra_count = self.n

        for _ in range(extra_count):
            s = GameServer()
            s.start()
            self.owned_servers.append(s)
            self.q.put(s)

        self.ex = ThreadPoolExecutor(max_workers=self.n)
        self.games = 0

    def _game(self, weights, opp):
        s = self.q.get()
        try:
            try:
                return s.run_game(weights, opp)
            except Exception:
                try:
                    s.stop()
                    s.start()
                    return s.run_game(weights, opp)
                except Exception:
                    return -1, MAX_CYCLES
        finally:
            self.q.put(s)

    def evaluate_many(self, weight_list, opp):
        """Evaluate a batch of candidate weights against an opponent (K_GAMES each)."""
        futs = [[self.ex.submit(self._game, w, opp) for _ in range(K_GAMES)]
                for w in weight_list]
        out = []
        for fs in futs:
            score = 0.0
            for f in fs:
                winner, _ = f.result()
                score += 1.0 if winner == 0 else (DRAW_SCORE if winner == -1 else 0.0)
            out.append(score / K_GAMES)
        self.games += len(weight_list) * K_GAMES
        return out

    def validate(self, weights):
        """Validate candidate weights against all 7 opponents (VALIDATION_GAMES each)."""
        futs = {o: [self.ex.submit(self._game, weights, o) for _ in range(VALIDATION_GAMES)]
                for o in OPPONENT_ROTATION}
        per_opp = {}
        total_wins = 0
        for o, fs in futs.items():
            wins = sum(1 for f in fs if f.result()[0] == 0)
            per_opp[o] = wins / VALIDATION_GAMES * 100.0
            total_wins += wins
        overall_wr = total_wins / (VALIDATION_GAMES * N_OPP) * 100.0
        return overall_wr, per_opp

    def close(self):
        self.ex.shutdown(wait=True)
        for s in self.owned_servers:
            try:
                s.stop()
            except Exception:
                pass


# ===============================================================
#  UPGRADED TRAJECTORY MECHANISM: BuildFingerprint & Optimal Transport
# ===============================================================
def _sse(X, idx):
    p = X[idx]
    return float(((p - p.mean(axis=0)) ** 2).sum())


def _best_split(X, idx):
    """2-means split of cluster `idx`."""
    if len(idx) < 2:
        return None
    p = X[idx]
    c = p.mean(axis=0)
    sse_parent = float(((p - c) ** 2).sum())
    if sse_parent < 1e-12:
        return None

    a = int(np.argmax(((p - c) ** 2).sum(axis=1)))
    b = int(np.argmax(((p - p[a]) ** 2).sum(axis=1)))
    if a == b:
        return None

    cen = np.stack([p[a], p[b]])
    lab = None
    for _ in range(25):
        dist = ((p[:, None, :] - cen[None, :, :]) ** 2).sum(axis=2)
        new = dist.argmin(axis=1)
        if lab is not None and (new == lab).all():
            break
        lab = new
        for k in (0, 1):
            mask = (lab == k)
            if mask.any():
                cen[k] = p[mask].mean(axis=0)

    left = [idx[i] for i in range(len(idx)) if lab[i] == 0]
    right = [idx[i] for i in range(len(idx)) if lab[i] == 1]
    if not left or not right:
        return None

    dE = sse_parent - _sse(X, left) - _sse(X, right)
    if dE <= 0.0:
        return None
    return dE, left, right


def build_subpop_fingerprint(params_list, fits, dim):
    """
    Builds the environmental modes fingerprint F_t = {(theta_m, w_m, g_m)}
    for a sub-population of dimension `dim`.
    """
    n = len(params_list)
    sqrt_dim = math.sqrt(dim)
    X = np.asarray(params_list, dtype=float) / sqrt_dim
    clusters = [list(range(n))]
    splits = {0: _best_split(X, clusters[0])}

    while True:
        M = len(clusters)
        tau_M = TAU0 * math.exp(LAM * (M - 1))
        candidates = [(s[0], i) for i, s in splits.items() if s is not None]
        if not candidates:
            break
        dE, best_i = max(candidates)
        if dE <= tau_M:
            break
        _, left, right = splits[best_i]
        clusters[best_i] = left
        clusters.append(right)
        splits[best_i] = _best_split(X, left)
        splits[len(clusters) - 1] = _best_split(X, right)

    reps, counts, gs = [], [], []
    for idx in clusters:
        centre = X[idx].mean(axis=0)
        near = idx[int(np.argmin(((X[idx] - centre) ** 2).sum(axis=1)))]
        reps.append([clamp01(x) for x in params_list[near]])
        counts.append(len(idx))
        gs.append(fits[near])

    return {
        "theta": np.asarray(reps),
        "count": counts,
        "w": [c / n for c in counts],
        "g": gs,
        "dim": dim,
        "uid": next_fp_uid(),
    }


def _hungarian(C):
    """Pure-Python Hungarian min-cost assignment fallback."""
    n = len(C)
    INF = float("inf")
    u, v = [0.0] * (n + 1), [0.0] * (n + 1)
    p, way = [0] * (n + 1), [0] * (n + 1)
    for i in range(1, n + 1):
        p[0] = i
        j0 = 0
        minv = [INF] * (n + 1)
        used = [False] * (n + 1)
        while True:
            used[j0] = True
            i0, delta, j1 = p[j0], INF, 0
            for j in range(1, n + 1):
                if not used[j]:
                    cur = C[i0 - 1][j - 1] - u[i0] - v[j]
                    if cur < minv[j]:
                        minv[j], way[j] = cur, j0
                    if minv[j] < delta:
                        delta, j1 = minv[j], j
            for j in range(n + 1):
                if used[j]:
                    u[p[j]] += delta
                    v[j] -= delta
                else:
                    minv[j] -= delta
            j0 = j1
            if p[j0] == 0:
                break
        while True:
            j1 = way[j0]
            p[j0] = p[j1]
            j0 = j1
            if j0 == 0:
                break
    return sum(C[p[j] - 1][j - 1] for j in range(1, n + 1))


def d_env_subpop(FA, FB):
    """Optimal transport distance between two sub-population fingerprints."""
    if FA.get("uid") == FB.get("uid"):
        return 0.0

    dim = FA["dim"]
    sqrt_dim = math.sqrt(dim)
    ia = np.repeat(np.arange(len(FA["count"])), FA["count"])
    ib = np.repeat(np.arange(len(FB["count"])), FB["count"])

    A, B = FA["theta"][ia], FB["theta"][ib]
    C = np.linalg.norm(A[:, None, :] - B[None, :, :], axis=2) / sqrt_dim

    if _lsa is not None:
        r, c = _lsa(C)
        return float(C[r, c].sum() / len(ia))
    return _hungarian(C.tolist()) / len(ia)


# ===============================================================
#  SUB-AGENT CLASSES: BASELINE vs UPGRADED TRAJECTORY
# ===============================================================
class BaselineSubAgent:
    """Sub-agent for Baseline CCE-TDCCEA (matching COMPARE_ALL.py)."""

    def __init__(self, dim, params=None, env_offset=0):
        self.dim        = dim
        self.env_offset = env_offset
        self.params     = ([random.random() for _ in range(dim)] if params is None
                           else [clamp01(x) for x in params])
        self.H_i  = []
        self.e_i  = 0.0
        self.mu_i = 0.0

    def update_history(self, f, env):
        self.H_i.append((f, list(env)))
        if len(self.H_i) > TRAJ_WINDOW:
            self.H_i.pop(0)

    def compute_trajectory(self, env):
        if not self.H_i:
            self.e_i = 0.0
            return
        sqrt_dim = math.sqrt(self.dim)
        off = self.env_offset
        ws = [max(0.0, 1.0 - math.dist(env[off : off + self.dim], pe[off : off + self.dim]) / sqrt_dim)
              if len(pe) >= off + self.dim else 1.0 for _, pe in self.H_i]
        tot = sum(ws)
        ws = [w / tot for w in ws] if tot > 0 else [1.0 / len(ws)] * len(ws)
        self.e_i = sum(w * self.H_i[k][0] for k, w in enumerate(ws))

    def compute_mutation(self):
        self.mu_i = MU_BASE * (1.0 - self.e_i) + EPS

    def mutate(self):
        return BaselineSubAgent(
            self.dim,
            [clamp01(p + random.gauss(0, self.mu_i)) for p in self.params],
            env_offset=self.env_offset,
        )

    def snapshot(self):
        c = BaselineSubAgent(self.dim, self.params, env_offset=self.env_offset)
        c.e_i, c.H_i, c.mu_i = self.e_i, self.H_i[:], self.mu_i
        return c


class UpgradedSubAgent:
    """Sub-agent with clustered strategic modes and optimal-transport trajectory memory."""

    def __init__(self, dim, params=None, env_offset=0):
        self.dim        = dim
        self.env_offset = env_offset
        self.params     = ([random.random() for _ in range(dim)] if params is None
                           else [clamp01(x) for x in params])
        self.H_i  = []
        self.e_i  = 0.0
        self.mu_i = 0.0

    def update_history(self, f, F):
        self.H_i.append((f, F))
        if len(self.H_i) > TRAJ_WINDOW:
            self.H_i.pop(0)

    def compute_trajectory(self, Ft, dist_cache):
        alphas = []
        for _, Fk in self.H_i:
            key = (Ft.get("uid"), Fk.get("uid"))
            if key not in dist_cache:
                dist_cache[key] = d_env_subpop(Ft, Fk)
            alphas.append(max(0.0, 1.0 - dist_cache[key] / D_MAX))
        tot = sum(alphas)
        alphas = [a / tot for a in alphas] if tot > 0 else [1.0 / len(alphas)] * len(alphas)
        self.e_i = sum(a * self.H_i[k][0] for k, a in enumerate(alphas))

    def compute_mutation(self):
        self.mu_i = MU_BASE * (1.0 - self.e_i) + EPS

    def mutate(self):
        return UpgradedSubAgent(
            self.dim,
            [clamp01(p + random.gauss(0, self.mu_i)) for p in self.params],
            env_offset=self.env_offset,
        )

    def snapshot(self):
        c = UpgradedSubAgent(self.dim, self.params, env_offset=self.env_offset)
        c.e_i, c.H_i, c.mu_i = self.e_i, self.H_i[:], self.mu_i
        return c


# ===============================================================
#  SUB-POPULATION SURVIVAL & REPRODUCTION (COMPARE_ALL.py logic)
# ===============================================================
def subpop_survival_step(pop, archive):
    """
    Sub-population survival:
      1. Redundancy filter: S(i, j) = 1 - d / sqrt(dim) > delta_s
      2. Weak-agent filter: e_i < delta_f and |H_i| >= 3
      3. Minimum floor: keep at least MIN_KEEP survivors (sorted by e_i)
      4. Elite archive update: top 3 from full pop inserted, cap A_max
      5. Tournament breeding: k=3 from survivors U archive
    """
    dim = pop[0].dim
    sqrt_dim = math.sqrt(dim)

    # 1. Redundancy removal
    ranked = sorted(pop, key=lambda a: a.e_i, reverse=True)
    survivors = []
    for a in ranked:
        if not any(1.0 - math.dist(a.params, u.params) / sqrt_dim > DELTA_S
                   for u in survivors):
            survivors.append(a)

    # 2. Weak-agent removal
    final = [a for a in survivors if not (a.e_i < DELTA_F and len(a.H_i) >= 3)]

    # 3. Minimum floor backfill (sorted by e_i descending)
    while len(final) < MIN_KEEP and len(survivors) > len(final):
        for a in sorted(survivors, key=lambda a: a.e_i, reverse=True):
            if a not in final:
                final.append(a)
                break
    final = final[:N_POP]

    # 4. Elite archive: top 3 from evaluated pop (matches COMPARE_ALL.py)
    for top in sorted(pop, key=lambda a: a.e_i, reverse=True)[:3]:
        archive.append(top.snapshot())
    archive.sort(key=lambda a: a.e_i, reverse=True)
    del archive[A_MAX:]

    # 5. Tournament reproduction
    breeding_pool = final + archive
    nxt = final[:]
    while len(nxt) < N_POP:
        candidates = random.sample(breeding_pool, min(3, len(breeding_pool)))
        parent = max(candidates, key=lambda a: a.e_i)
        nxt.append(parent.mutate())

    return nxt, archive


# ===============================================================
#  ALGORITHM 1: BASELINE CCE-TDCCEA (Centroid Trajectory)
# ===============================================================
def run_cce_tdccea_baseline(server=None, n_workers=None, return_history=False):
    print(f"\n{'='*72}")
    print("  RUNNING: Base CCE-TDCCEA (Macro/Micro Decomposition + Centroid Trajectory)")
    print(f"{'='*72}")
    t_start = time.time()

    if isinstance(server, ServerPool):
        pool = server
        owns_pool = False
    else:
        pool = ServerPool(base_server=server, n_workers=n_workers)
        owns_pool = True

    print(f"  Parallel game workers: {pool.n}")
    dim_macro = len(MACRO_INDICES)
    dim_micro = len(MICRO_INDICES)

    try:
        pop_a = [BaselineSubAgent(dim_macro, env_offset=0) for _ in range(N_POP)]
        pop_b = [BaselineSubAgent(dim_micro, env_offset=dim_macro) for _ in range(N_POP)]
        archive_a, archive_b = [], []
        top_a, top_b = pop_a[:TOP_PARTNERS], pop_b[:TOP_PARTNERS]
        convergence = []

        overall_best_weights = combine_subagents(pop_a[0], pop_b[0])
        overall_best_fit = -1.0

        for gen in range(1, TOTAL_GENERATIONS + 1):
            opp = OPPONENT_ROTATION[(gen - 1) % N_OPP]

            # 1. Parallel Cooperative Evaluation
            # Pop A: evaluate with top_b partners
            combos_a = [combine_subagents(a, pb) for a in pop_a for pb in top_b]
            scores_a = pool.evaluate_many(combos_a, opp)
            k_b = len(top_b)
            fits_a = [sum(scores_a[i * k_b : (i + 1) * k_b]) / k_b for i in range(N_POP)]

            # Pop B: evaluate with top_a partners
            combos_b = [combine_subagents(pa, b) for b in pop_b for pa in top_a]
            scores_b = pool.evaluate_many(combos_b, opp)
            k_a = len(top_a)
            fits_b = [sum(scores_b[j * k_a : (j + 1) * k_a]) / k_a for j in range(N_POP)]

            # 2. Ecosystem Centroids (matches COMPARE_ALL.py env = ca + cb)
            ca = [sum(a.params[j] for a in pop_a) / len(pop_a) for j in range(dim_macro)]
            cb = [sum(b.params[j] for b in pop_b) / len(pop_b) for j in range(dim_micro)]
            env = ca + cb

            # 3. Trajectory scoring with sub-population centroids
            for a, f in zip(pop_a, fits_a):
                a.update_history(f, env)
                a.compute_trajectory(env)
                a.compute_mutation()

            for b, f in zip(pop_b, fits_b):
                b.update_history(f, env)
                b.compute_trajectory(env)
                b.compute_mutation()

            # 4. Best agents & Top partners
            best_a = max(pop_a, key=lambda a: a.e_i)
            best_b = max(pop_b, key=lambda b: b.e_i)
            top_a = sorted(pop_a, key=lambda a: a.e_i, reverse=True)[:TOP_PARTNERS]
            top_b = sorted(pop_b, key=lambda b: b.e_i, reverse=True)[:TOP_PARTNERS]

            # Combined evaluation of co-adapted pair
            current_best = combine_subagents(best_a, best_b)
            cmb_fit = pool.evaluate_many([current_best], opp)[0]
            convergence.append(cmb_fit)

            if cmb_fit > overall_best_fit:
                overall_best_fit = cmb_fit
                overall_best_weights = list(current_best)

            print(f"  Gen {gen:2d}/{TOTAL_GENERATIONS} | {opp:25s} | "
                  f"A_e={best_a.e_i:.3f} B_e={best_b.e_i:.3f} Combined={cmb_fit:.3f} | games={pool.games}")

            # 5. Sub-population survival & reproduction
            pop_a, archive_a = subpop_survival_step(pop_a, archive_a)
            pop_b, archive_b = subpop_survival_step(pop_b, archive_b)

        elapsed = time.time() - t_start

        if return_history:
            return overall_best_weights, elapsed, convergence
        return overall_best_weights, elapsed

    finally:
        if owns_pool:
            pool.close()


# ===============================================================
#  ALGORITHM 2: CCE-TDCCEA WITH UPGRADED TRAJECTORY MECHANISM
# ===============================================================
def run_cce_tdccea_upgraded_traj(server=None, n_workers=None, return_history=False):
    print(f"\n{'='*72}")
    print("  RUNNING: CCE-TDCCEA (Upgraded Trajectory: Strategic Modes + Optimal Transport)")
    print(f"{'='*72}")
    t_start = time.time()

    if isinstance(server, ServerPool):
        pool = server
        owns_pool = False
    else:
        pool = ServerPool(base_server=server, n_workers=n_workers)
        owns_pool = True

    print(f"  Parallel game workers: {pool.n}")
    dim_macro = len(MACRO_INDICES)
    dim_micro = len(MICRO_INDICES)

    try:
        pop_a = [UpgradedSubAgent(dim_macro, env_offset=0) for _ in range(N_POP)]
        pop_b = [UpgradedSubAgent(dim_micro, env_offset=dim_macro) for _ in range(N_POP)]
        archive_a, archive_b = [], []
        top_a, top_b = pop_a[:TOP_PARTNERS], pop_b[:TOP_PARTNERS]
        convergence = []
        modes_history = []

        overall_best_weights = combine_subagents(pop_a[0], pop_b[0])
        overall_best_fit = -1.0

        for gen in range(1, TOTAL_GENERATIONS + 1):
            opp = OPPONENT_ROTATION[(gen - 1) % N_OPP]

            # 1. Parallel Cooperative Evaluation
            # Pop A: evaluate with top_b partners
            combos_a = [combine_subagents(a, pb) for a in pop_a for pb in top_b]
            scores_a = pool.evaluate_many(combos_a, opp)
            k_b = len(top_b)
            fits_a = [sum(scores_a[i * k_b : (i + 1) * k_b]) / k_b for i in range(N_POP)]

            # Pop B: evaluate with top_a partners
            combos_b = [combine_subagents(pa, b) for b in pop_b for pa in top_a]
            scores_b = pool.evaluate_many(combos_b, opp)
            k_a = len(top_a)
            fits_b = [sum(scores_b[j * k_a : (j + 1) * k_a]) / k_a for j in range(N_POP)]

            # 2. Clustered Strategic Modes Fingerprint per sub-population (Algorithm 1)
            Fa = build_subpop_fingerprint([a.params for a in pop_a], fits_a, dim_macro)
            Fb = build_subpop_fingerprint([b.params for b in pop_b], fits_b, dim_micro)

            modes_history.append({
                "gen": gen,
                "modes_a": len(Fa["count"]),
                "modes_b": len(Fb["count"]),
                "sizes_a": sorted(Fa["count"], reverse=True),
                "sizes_b": sorted(Fb["count"], reverse=True),
            })

            # 3. Trajectory scoring with Optimal-Transport distance
            cache_a = {}
            for a, f in zip(pop_a, fits_a):
                a.update_history(f, Fa)
            for a in pop_a:
                a.compute_trajectory(Fa, cache_a)
                a.compute_mutation()

            cache_b = {}
            for b, f in zip(pop_b, fits_b):
                b.update_history(f, Fb)
            for b in pop_b:
                b.compute_trajectory(Fb, cache_b)
                b.compute_mutation()

            # 4. Best agents & Top partners
            best_a = max(pop_a, key=lambda a: a.e_i)
            best_b = max(pop_b, key=lambda b: b.e_i)
            top_a = sorted(pop_a, key=lambda a: a.e_i, reverse=True)[:TOP_PARTNERS]
            top_b = sorted(pop_b, key=lambda b: b.e_i, reverse=True)[:TOP_PARTNERS]

            # Combined evaluation of co-adapted pair
            current_best = combine_subagents(best_a, best_b)
            cmb_fit = pool.evaluate_many([current_best], opp)[0]
            convergence.append(cmb_fit)

            if cmb_fit > overall_best_fit:
                overall_best_fit = cmb_fit
                overall_best_weights = list(current_best)

            print(f"  Gen {gen:2d}/{TOTAL_GENERATIONS} | {opp:25s} | "
                  f"A_e={best_a.e_i:.3f} B_e={best_b.e_i:.3f} Combined={cmb_fit:.3f} | "
                  f"modes A={len(Fa['count'])}, B={len(Fb['count'])} | games={pool.games}")

            # 5. Sub-population survival & reproduction
            pop_a, archive_a = subpop_survival_step(pop_a, archive_a)
            pop_b, archive_b = subpop_survival_step(pop_b, archive_b)

        elapsed = time.time() - t_start

        if return_history:
            return overall_best_weights, elapsed, convergence, modes_history
        return overall_best_weights, elapsed

    finally:
        if owns_pool:
            pool.close()


# ===============================================================
#  WILCOXON SIGNED-RANK TEST (Two-Sided)
# ===============================================================
def wilcoxon_test(x, y):
    diffs = [xi - yi for xi, yi in zip(x, y) if xi != yi]
    if len(diffs) < 3:
        return None
    n = len(diffs)
    abs_d = sorted(enumerate(abs(d) for d in diffs), key=lambda t: t[1])
    ranks = [0.0] * n
    i = 0
    while i < n:
        j = i
        while j < n - 1 and abs_d[j][1] == abs_d[j + 1][1]:
            j += 1
        avg_rank = (i + j) / 2.0 + 1
        for k in range(i, j + 1):
            ranks[abs_d[k][0]] = avg_rank
        i = j + 1
    w_plus  = sum(ranks[i] for i, d in enumerate(diffs) if d > 0)
    w_minus = sum(ranks[i] for i, d in enumerate(diffs) if d < 0)
    w_stat  = min(w_plus, w_minus)
    mu_w    = n * (n + 1) / 4.0
    sig_w   = math.sqrt(n * (n + 1) * (2 * n + 1) / 24.0)
    if sig_w == 0:
        return None
    z = (w_stat - mu_w) / sig_w
    p = 2.0 * (1.0 - 0.5 * (1.0 + math.erf(abs(z) / math.sqrt(2.0))))
    return p


# ===============================================================
#  MAIN BENCHMARK & COMPARATIVE EVALUATION
# ===============================================================
def main():
    nw = int(sys.argv[1]) if len(sys.argv) > 1 and sys.argv[1].isdigit() else default_workers()

    print("\n" + "#" * 72)
    print("  CCE-TDCCEA: BASELINE vs UPGRADED TRAJECTORY MECHANISM")
    print("  Direct Comparison: Centroid Trajectory vs Clustered Modes + Optimal Transport")
    print("#" * 72)
    print(f"  Pop A (Macro)={N_POP} | Pop B (Micro)={N_POP} | Gens={TOTAL_GENERATIONS}")
    print(f"  Collaborating Partners: Top-{TOP_PARTNERS} | K_games={K_GAMES} | Validation={VALIDATION_GAMES}/bot")
    print(f"  Map: {MAP_PATH} (32x32)")
    print(f"  Parallel game workers: {nw} (of {os.cpu_count()} CPU cores)")
    print(f"  Optimal transport solver: {'scipy linear_sum_assignment' if _lsa is not None else 'pure-Python Hungarian'}")
    print()

    pool = ServerPool(n_workers=nw)
    results = []

    try:
        # -------------------------------------------------------------
        # 1. Base CCE-TDCCEA (Centroid Trajectory)
        # -------------------------------------------------------------
        best_base, time_base, conv_base = run_cce_tdccea_baseline(pool, return_history=True)
        print("\n  Validating Base CCE-TDCCEA against all 7 bots...")
        wr_base, opp_base = pool.validate(best_base)
        print(f"  Base CCE-TDCCEA: {wr_base:.1f}% overall in {time_base:.1f}s")
        results.append({
            "name": "CCE-TDCCEA (Baseline)",
            "key": "base_cce",
            "time": time_base,
            "wr": wr_base,
            "per_opp": opp_base,
            "weights": best_base,
            "convergence": conv_base,
            "modes_history": None,
        })

        # -------------------------------------------------------------
        # 2. CCE-TDCCEA (Upgraded Trajectory)
        # -------------------------------------------------------------
        best_up, time_up, conv_up, modes_up = run_cce_tdccea_upgraded_traj(pool, return_history=True)
        print("\n  Validating Upgraded CCE-TDCCEA against all 7 bots...")
        wr_up, opp_up = pool.validate(best_up)
        print(f"  Upgraded CCE-TDCCEA: {wr_up:.1f}% overall in {time_up:.1f}s")
        results.append({
            "name": "CCE-TDCCEA (Upgraded Traj)",
            "key": "upgraded_cce",
            "time": time_up,
            "wr": wr_up,
            "per_opp": opp_up,
            "weights": best_up,
            "convergence": conv_up,
            "modes_history": modes_up,
        })

    finally:
        pool.close()

    # ═══════════════════════════════════════════════════════════════
    #  FINAL RANKING TABLE
    # ═══════════════════════════════════════════════════════════════
    ranked = sorted(results, key=lambda r: r["wr"], reverse=True)
    medals = ["1st", "2nd"]
    print(f"\n{'='*72}")
    print("  FINAL RANKING — Win Rate vs 7 microRTS Bots")
    print(f"{'='*72}")
    print(f"  {'Rank':<6s} {'Algorithm':<28s} {'Win Rate':>10s} {'Time':>9s}")
    print(f"  {'_'*6} {'_'*28} {'_'*10} {'_'*9}")
    for i, r in enumerate(ranked):
        tag = "  ◄ Proposed (Modes + OT)" if r["key"] == "upgraded_cce" else "  ◄ Baseline (Centroid)"
        print(f"  {medals[i]:<6s} {r['name']:<28s} {r['wr']:>9.1f}%  {r['time']:>8.1f}s{tag}")

    # ═══════════════════════════════════════════════════════════════
    #  PER-BOT BREAKDOWN TABLE
    # ═══════════════════════════════════════════════════════════════
    print(f"\n{'='*72}")
    print("  WIN RATE PER BOT (%)")
    print(f"{'='*72}")
    header = f"  {'Bot':<24s}"
    for r in ranked:
        header += f" {r['name'][:20]:>20s}"
    print(header)
    print("  " + "-" * (24 + 21 * len(ranked)))
    for opp in OPPONENT_ROTATION:
        row = f"  {opp:<24s}"
        rates = [r["per_opp"].get(opp, 0.0) for r in ranked]
        mx = max(rates)
        for r in ranked:
            w = r["per_opp"].get(opp, 0.0)
            mk = "*" if w == mx and rates.count(mx) == 1 else " "
            row += f" {w:>18.0f}%{mk}"
        print(row)
    print("  " + "-" * (24 + 21 * len(ranked)))
    row = f"  {'OVERALL':<24s}"
    for r in ranked:
        row += f" {r['wr']:>19.1f}% "
    print(row)

    # ═══════════════════════════════════════════════════════════════
    #  WILCOXON SIGNED-RANK TEST
    # ═══════════════════════════════════════════════════════════════
    print(f"\n{'='*72}")
    print("  STATISTICAL SIGNIFICANCE — Wilcoxon Signed-Rank Test")
    print("  (CCE-TDCCEA Upgraded Trajectory vs Baseline, 7 paired opponent samples)")
    print(f"{'='*72}")

    up_res = next((r for r in results if r["key"] == "upgraded_cce"), None)
    base_res = next((r for r in results if r["key"] == "base_cce"), None)

    w_out = [f"\nRun Date: {time.strftime('%Y-%m-%d %H:%M:%S')}"]
    if up_res and base_res:
        up_rates = [up_res["per_opp"].get(o, 0.0) for o in OPPONENT_ROTATION]
        base_rates = [base_res["per_opp"].get(o, 0.0) for o in OPPONENT_ROTATION]
        p = wilcoxon_test(up_rates, base_rates)
        if p is None:
            line = "  p-value: N/A (fewer than 3 non-tied differences)"
        else:
            sig = "YES (p < 0.05)" if p < 0.05 else "NO (p >= 0.05)"
            dirn = "UPGRADED TRAJECTORY BETTER" if sum(up_rates) > sum(base_rates) else "BASELINE BETTER"
            line = f"  p-value: {p:.4f}  |  Significant: {sig}  |  Direction: {dirn}"
        print(f"\n{line}")
        w_out.append(line)

        print(f"\n  {'Bot':<24s} {'Baseline':>12s} {'Upgraded':>14s} {'Delta':>8s}")
        print(f"  {'_'*24} {'_'*12} {'_'*14} {'_'*8}")
        for opp in OPPONENT_ROTATION:
            b_wr = base_res["per_opp"].get(opp, 0.0)
            u_wr = up_res["per_opp"].get(opp, 0.0)
            delta = u_wr - b_wr
            sign = "+" if delta >= 0 else ""
            print(f"  {opp:<24s} {b_wr:>11.0f}% {u_wr:>13.0f}% {sign}{delta:>6.0f}%")

    wilc_path = os.path.join(SCRIPT_DIR, "cce_trajectory_wilcoxon.txt")
    with open(wilc_path, "a", encoding="utf-8") as f:
        f.write("\n".join(w_out) + "\n\n")

    # ═══════════════════════════════════════════════════════════════
    #  ABLATION ANALYSIS
    # ═══════════════════════════════════════════════════════════════
    print(f"\n{'='*72}")
    print("  ABLATION ANALYSIS: Upgraded Trajectory vs Baseline Centroid Trajectory")
    print(f"{'='*72}")
    if up_res and base_res:
        diff_wr = up_res["wr"] - base_res["wr"]
        time_oh = up_res["time"] - base_res["time"]
        mode_counts_a = [m["modes_a"] for m in up_res["modes_history"]] if up_res["modes_history"] else []
        mode_counts_b = [m["modes_b"] for m in up_res["modes_history"]] if up_res["modes_history"] else []
        avg_a = sum(mode_counts_a) / len(mode_counts_a) if mode_counts_a else 1.0
        avg_b = sum(mode_counts_b) / len(mode_counts_b) if mode_counts_b else 1.0

        lines = [
            "  Trajectory Mechanism Impact on CCE-TDCCEA:",
            f"    CCE-TDCCEA (Upgraded Modes + Optimal Transport): {up_res['wr']:.1f}% win rate",
            f"    CCE-TDCCEA (Baseline Single Centroid):           {base_res['wr']:.1f}% win rate",
            f"    Net Gain from Upgraded Trajectory Mechanism:     {'+' if diff_wr >= 0 else ''}{diff_wr:.1f}%",
            "  Computational Performance:",
            f"    Upgraded Trajectory Runtime: {up_res['time']:.1f}s",
            f"    Baseline Centroid Runtime:   {base_res['time']:.1f}s",
            f"    Clustering & Transport Overhead: {time_oh:+.1f}s",
            "  Sub-Population Strategic Modes:",
            f"    Macro Sub-Population (8 weights): avg {avg_a:.2f} modes/gen",
            f"    Micro Sub-Population (12 weights): avg {avg_b:.2f} modes/gen",
        ]
        for l in lines:
            print(f"  {l}")

        ablation_path = os.path.join(SCRIPT_DIR, "cce_trajectory_ablation.txt")
        with open(ablation_path, "a", encoding="utf-8") as f:
            f.write(time.strftime("%Y-%m-%d %H:%M:%S") + "\n" + "\n".join(lines) + "\n\n")

    # ═══════════════════════════════════════════════════════════════
    #  CONVERGENCE TABLE
    # ═══════════════════════════════════════════════════════════════
    print(f"\n{'='*72}")
    print("  CONVERGENCE COMPARISON (Best Combined Fitness per Generation)")
    print(f"{'='*72}")
    print(f"  {'Gen':>4s}  {'Baseline CCE':>18s}  {'Upgraded Traj CCE':>20s}")
    print(f"  {'_'*4}  {'_'*18}  {'_'*20}")
    max_gens = min(len(results[0]["convergence"]), len(results[1]["convergence"]))
    for g in range(max_gens):
        b_val = results[0]["convergence"][g]
        u_val = results[1]["convergence"][g]
        print(f"  {g+1:4d}  {b_val:>18.3f}  {u_val:>20.3f}")

    # ═══════════════════════════════════════════════════════════════
    #  JSON & EXCEL EXPORTS
    # ═══════════════════════════════════════════════════════════════
    json_path = os.path.join(SCRIPT_DIR, "cce_trajectory_results.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump([{
            "name": r["name"],
            "key": r["key"],
            "wr": r["wr"],
            "per_opp": r["per_opp"],
            "time": r["time"],
            "weights": r["weights"],
            "convergence": r["convergence"],
            "modes_history": r["modes_history"],
        } for r in results], f, indent=2)
    print(f"\n  JSON saved -> {json_path}")

    if _OPENPYXL_AVAILABLE:
        xlsx_path = os.path.join(SCRIPT_DIR, "cce_trajectory_comparison.xlsx")
        algo_names = [r["name"] for r in results]
        algo_wr    = {r["name"]: round(r["wr"], 2) for r in results}
        algo_time  = {r["name"]: round(r["time"], 1) for r in results}

        if not os.path.exists(xlsx_path):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Summary"
            ws.append(["Algorithm", "Run1_WinRate", "Run1_Time"])
            for a in algo_names:
                ws.append([a, algo_wr[a], algo_time[a]])
            wb.save(xlsx_path)
            run_col = "Run1"
        else:
            wb = openpyxl.load_workbook(xlsx_path)
            ws = wb["Summary"] if "Summary" in wb.sheetnames else wb.active
            header_row = [cell.value for cell in ws[1]]
            run_numbers = [
                int(h.split("_")[0][3:]) for h in header_row
                if isinstance(h, str) and h.lower().startswith("run") and "_" in h and h.split("_")[0][3:].isdigit()
            ]
            next_run = (max(run_numbers) + 1) if run_numbers else 1
            run_col = f"Run{next_run}"
            time_col = f"Time{next_run}"
            r_idx = len(header_row) + 1
            t_idx = len(header_row) + 2
            ws.cell(row=1, column=r_idx, value=f"{run_col}_WinRate")
            ws.cell(row=1, column=t_idx, value=f"{run_col}_{time_col}")

            row_map = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1)}
            for a in algo_names:
                row = row_map.get(a) or (ws.max_row + 1)
                ws.cell(row=row, column=1, value=a)
                ws.cell(row=row, column=r_idx, value=algo_wr[a])
                ws.cell(row=row, column=t_idx, value=algo_time[a])
            wb.save(xlsx_path)

        print(f"  Excel updated -> {xlsx_path} (Run: {run_col})")
    else:
        print("  [SKIP] openpyxl not installed — Excel export skipped.")

    print(f"\n{'='*72}\n  BENCHMARK COMPLETE!\n{'='*72}\n")


if __name__ == "__main__":
    main()
